#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка contacts_russia.xlsx + 3 CSV (UTF-8 BOM) из leads.json."""
import json, csv, os, datetime, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

COLS = ["ID","Категория","Название мастерской / имя мастера","Никнейм","Город","Регион",
        "Что производит","Признаки собственного производства","Активность",
        "Последняя обнаруженная активность","Основной способ связи","Основной контакт",
        "VK","Telegram","Instagram / другая соцсеть","Сайт","Email","Дополнительные контакты",
        "Основной источник","Дополнительный источник","Почему подходит нашей мастерской",
        "Потенциально интересующие детали / услуги","Коммерческая релевантность 1-5",
        "Качество лида A/B/C","Комментарий исследователя","Дата проверки"]
URLCOLS = {"Основной контакт","VK","Telegram","Instagram / другая соцсеть","Сайт",
           "Основной источник","Дополнительный источник"}
WIDTHS = {"ID":6,"Категория":22,"Название мастерской / имя мастера":30,"Никнейм":20,"Город":18,
          "Регион":22,"Что производит":38,"Признаки собственного производства":38,"Активность":16,
          "Последняя обнаруженная активность":24,"Основной способ связи":18,"Основной контакт":34,
          "VK":34,"Telegram":28,"Instagram / другая соцсеть":30,"Сайт":32,"Email":28,
          "Дополнительные контакты":30,"Основной источник":40,"Дополнительный источник":36,
          "Почему подходит нашей мастерской":46,"Потенциально интересующие детали / услуги":42,
          "Коммерческая релевантность 1-5":14,"Качество лида A/B/C":14,
          "Комментарий исследователя":46,"Дата проверки":13}

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
LINK_FONT = Font(color="0563C1", underline="single")

def is_url(v): return isinstance(v,str) and v.startswith(("http://","https://"))

def sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    ws.append(COLS)
    for c in range(1,len(COLS)+1):
        cell = ws.cell(1,c); cell.fill=HDR_FILL; cell.font=HDR_FONT
        cell.alignment=Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 46
    for r in rows:
        ws.append([r.get(c,"") for c in COLS])
    for ci,cname in enumerate(COLS,1):
        ws.column_dimensions[get_column_letter(ci)].width = WIDTHS.get(cname,20)
        if cname in URLCOLS:
            for ri in range(2, ws.max_row+1):
                cell = ws.cell(ri,ci)
                if is_url(cell.value):
                    cell.hyperlink = cell.value; cell.font = LINK_FONT
    for ri in range(2, ws.max_row+1):
        for ci in range(1,len(COLS)+1):
            ws.cell(ri,ci).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
    return ws

def summary(wb, fur, cos, unv):
    ws = wb.create_sheet("Summary")
    allq = fur + cos
    def cnt(rows, key, val): return sum(1 for r in rows if str(r.get(key,"")).strip().upper()==val)
    def has(rows, key): return sum(1 for r in rows if str(r.get(key,"")).strip() not in ("","-","н/д"))
    rel = [float(r["Коммерческая релевантность 1-5"]) for r in allq
           if str(r.get("Коммерческая релевантность 1-5","")).replace(".","",1).isdigit()]
    block = [("ПОКАЗАТЕЛЬ","ЗНАЧЕНИЕ"),
        ("Fursuit-лидов", len(fur)), ("Cosplay-лидов", len(cos)),
        ("Требует проверки", len(unv)), ("ВСЕГО записей", len(fur)+len(cos)+len(unv)), ("",""),
        ("A-лидов", cnt(allq,"Качество лида A/B/C","A")),
        ("B-лидов", cnt(allq,"Качество лида A/B/C","B")),
        ("C-лидов", cnt(allq,"Качество лида A/B/C","C")), ("",""),
        ("Лидов с Telegram", has(allq,"Telegram")),
        ("Лидов с VK", has(allq,"VK")),
        ("Лидов с Email", has(allq,"Email")),
        ("Лидов с сайтом", has(allq,"Сайт")), ("",""),
        ("Средняя коммерческая релевантность", round(sum(rel)/len(rel),2) if rel else 0)]
    for row in block: ws.append(list(row))
    ws.append([]); ws.append(["РАСПРЕДЕЛЕНИЕ ПО РЕГИОНАМ","КОЛ-ВО"])
    for k,v in collections.Counter((r.get("Регион") or "не указан") for r in allq).most_common(): ws.append([k,v])
    ws.append([]); ws.append(["РАСПРЕДЕЛЕНИЕ ПО ТИПУ ПРОИЗВОДИТЕЛЕЙ","КОЛ-ВО"])
    for k,v in collections.Counter((r.get("Категория") or "не указан") for r in allq).most_common(): ws.append([k,v])
    ws.append([]); ws.append(["РАСПРЕДЕЛЕНИЕ ПО ГОРОДАМ (топ)","КОЛ-ВО"])
    for k,v in collections.Counter((r.get("Город") or "не указан") for r in allq).most_common(20): ws.append([k,v])
    for ri in range(1, ws.max_row+1):
        a = ws.cell(ri,1)
        if a.value and str(a.value).isupper() and len(str(a.value))>8:
            a.font = Font(bold=True, color="FFFFFF"); a.fill = HDR_FILL
            ws.cell(ri,2).font = Font(bold=True, color="FFFFFF"); ws.cell(ri,2).fill = HDR_FILL
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A2"
    return ws

def write_csv(path, rows):
    with open(path,"w",encoding="utf-8-sig",newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS, extrasaction="ignore"); w.writeheader()
        for r in rows: w.writerow({c: r.get(c,"") for c in COLS})

def main():
    data = json.load(open("leads.json",encoding="utf-8"))
    fur, cos, unv = data["fursuit"], data["cosplay"], data["unverified"]
    for group,pfx in ((fur,"F"),(cos,"C"),(unv,"U")):
        for i,r in enumerate(group,1): r["ID"] = f"{pfx}{i:03d}"
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb,"Fursuit",fur); sheet(wb,"Cosplay",cos); sheet(wb,"Требует проверки",unv)
    summary(wb,fur,cos,unv)
    wb.save("contacts_russia.xlsx")
    write_csv("fursuit_contacts_russia.csv",fur)
    write_csv("cosplay_contacts_russia.csv",cos)
    write_csv("unverified_contacts_russia.csv",unv)
    print(f"OK: Fursuit={len(fur)} Cosplay={len(cos)} Требует проверки={len(unv)}")

if __name__=="__main__": main()
